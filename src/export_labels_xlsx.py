"""Export data/labels/sample.csv to an Excel file for hand labelling.

The Excel file has a frozen header, readable column widths, highlighted label columns
and drop-down menus for my_category and my_confidence. No API calls.

Refuses to overwrite an existing Excel file (it may hold unsaved labels) unless --force is given.
"""

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from config import PROJECT_ROOT
from prompt import CATEGORIES, CONFIDENCE_LEVELS

SAMPLE_FILE = PROJECT_ROOT / "data" / "labels" / "sample.csv"
XLSX_FILE = PROJECT_ROOT / "data" / "labels" / "sample_labelling.xlsx"

COLUMN_WIDTHS = {
    "company_id": 12,
    "name": 40,
    "country": 18,
    "industry": 24,
    "size_group": 11,
    "split": 7,
    "my_category": 16,
    "my_confidence": 15,
    "my_note": 80,
}
LABEL_COLUMNS = ["my_category", "my_confidence", "my_note"]
LABEL_FILL = PatternFill(fill_type="solid", start_color="FFF2CC")  # pale yellow: type here


def add_dropdown(ws, column_letter: str, values: list[str], n_rows: int) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.error = "Choose one of: " + ", ".join(values)
    dv.errorTitle = "Invalid value"
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}2:{column_letter}{n_rows + 1}")


def write_workbook(sample: pd.DataFrame, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "labels"

    ws.append(list(sample.columns))
    for row in sample.itertuples(index=False):
        ws.append(list(row))

    letters = {}
    for cell in ws[1]:
        name = cell.value
        letters[name] = cell.column_letter
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = COLUMN_WIDTHS.get(name, 12)

    for name in ["name", "country", "industry"]:
        for cell in ws[letters[name]][1:]:
            cell.font = Font(bold=True)

    for name in LABEL_COLUMNS:
        for cell in ws[letters[name]][1:]:
            cell.fill = LABEL_FILL
    for cell in ws[letters["my_note"]][1:]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    add_dropdown(ws, letters["my_category"], CATEGORIES, len(sample))
    add_dropdown(ws, letters["my_confidence"], CONFIDENCE_LEVELS, len(sample))

    ws.freeze_panes = "A2"
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Export the labelling sample to Excel.")
    parser.add_argument("--output", default=XLSX_FILE)
    parser.add_argument("--force", action="store_true", help="Overwrite an existing Excel file.")
    args = parser.parse_args()

    output = Path(args.output)
    if output.exists() and not args.force:
        raise SystemExit(f"{output} already exists and may contain labels. Use --force to overwrite.")

    sample = pd.read_csv(SAMPLE_FILE, dtype=str, keep_default_na=False)
    write_workbook(sample, output)
    print(f"Wrote {len(sample)} rows to {output}")


if __name__ == "__main__":
    main()
